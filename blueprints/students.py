from flask import render_template, request, jsonify, session
from core.database import get_db_connection
from core.logger import get_logger
from blueprints.decorators import admin_required
from datetime import datetime

logger = get_logger(__name__)


def get_today_jalali():
    try:
        import jdatetime
        return jdatetime.date.today().strftime('%Y/%m/%d')
    except Exception:
        return datetime.now().strftime('%Y-%m-%d')


def get_current_school_id():
    """school_id کاربر فعلی رو برمی‌گردونه"""
    return session.get('school_id')


def students_page():
    return render_template('students.html')


# =====================================================
# لیست دانش‌آموزان
# =====================================================
def get_students():
    try:
        school_id = get_current_school_id()
        if not school_id:
            return jsonify({'success': False, 'error': 'مدرسه انتخاب نشده'})

        grade = request.args.get('grade', '').strip()
        class_name = request.args.get('class_name', '').strip()
        search = request.args.get('search', '').strip()

        query = "SELECT * FROM students WHERE is_active = 1 AND school_id = ?"
        params = [school_id]

        if grade:
            query += " AND grade = ?"
            params.append(grade)
        if class_name:
            query += " AND class_name = ?"
            params.append(class_name)
        if search:
            query += """ AND (
                first_name LIKE ? OR
                last_name LIKE ? OR
                national_code LIKE ? OR
                father_name LIKE ?
            )"""
            like = f"%{search}%"
            params.extend([like, like, like, like])

        query += " ORDER BY grade, class_name, last_name, first_name"

        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(query, params)
        rows = cursor.fetchall()
        conn.close()

        return jsonify({
            'success': True,
            'students': [dict(r) for r in rows],
            'count': len(rows)
        })
    except Exception as e:
        logger.error(f"خطا در لیست دانش‌آموزان: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500


# =====================================================
# افزودن دانش‌آموز
# =====================================================
@admin_required
def add_student():
    try:
        school_id = get_current_school_id()
        if not school_id:
            return jsonify({'success': False, 'error': 'مدرسه انتخاب نشده'})

        data = request.get_json()
        national_code = data.get('national_code', '').strip()
        first_name = data.get('first_name', '').strip()
        last_name = data.get('last_name', '').strip()
        father_name = data.get('father_name', '').strip()
        grade = data.get('grade', '').strip()
        class_name = data.get('class_name', '').strip()
        field = data.get('field', '').strip()
        phone1 = data.get('phone1', '').strip()
        phone2 = data.get('phone2', '').strip()

        if not national_code or not first_name or not last_name or not grade or not class_name:
            return jsonify({'success': False, 'error': 'فیلدهای اجباری را پر کنید'})

        if len(national_code) != 10 or not national_code.isdigit():
            return jsonify({'success': False, 'error': 'کد ملی باید ۱۰ رقم عددی باشد'})

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute(
            "SELECT id FROM students WHERE national_code = ? AND school_id = ?",
            (national_code, school_id)
        )
        if cursor.fetchone():
            conn.close()
            return jsonify({'success': False, 'error': 'این کد ملی قبلاً در این مدرسه ثبت شده'})

        cursor.execute('''
            INSERT INTO students
            (school_id, national_code, first_name, last_name, father_name, grade, class_name, field, phone1, phone2, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (school_id, national_code, first_name, last_name, father_name, grade, class_name, field, phone1, phone2, get_today_jalali()))

        conn.commit()
        student_id = cursor.lastrowid
        conn.close()

        return jsonify({'success': True, 'message': 'دانش‌آموز اضافه شد', 'id': student_id})
    except Exception as e:
        logger.error(f"خطا در افزودن دانش‌آموز: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500


# =====================================================
# ویرایش دانش‌آموز
# =====================================================
@admin_required
def edit_student(student_id):
    try:
        school_id = get_current_school_id()
        data = request.get_json()

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT id FROM students WHERE id = ? AND school_id = ?", (student_id, school_id))
        if not cursor.fetchone():
            conn.close()
            return jsonify({'success': False, 'error': 'دانش‌آموز یافت نشد'})

        cursor.execute('''
            UPDATE students SET
                first_name = ?, last_name = ?, father_name = ?,
                grade = ?, class_name = ?, field = ?,
                phone1 = ?, phone2 = ?
            WHERE id = ? AND school_id = ?
        ''', (
            data.get('first_name', '').strip(),
            data.get('last_name', '').strip(),
            data.get('father_name', '').strip(),
            data.get('grade', '').strip(),
            data.get('class_name', '').strip(),
            data.get('field', '').strip(),
            data.get('phone1', '').strip(),
            data.get('phone2', '').strip(),
            student_id, school_id
        ))

        conn.commit()
        conn.close()
        return jsonify({'success': True, 'message': 'تغییرات ذخیره شد'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


# =====================================================
# حذف دانش‌آموز
# =====================================================
@admin_required
def delete_student(student_id):
    try:
        school_id = get_current_school_id()
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE students SET is_active = 0 WHERE id = ? AND school_id = ?",
            (student_id, school_id)
        )
        conn.commit()
        conn.close()
        return jsonify({'success': True, 'message': 'دانش‌آموز حذف شد'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


# =====================================================
# لیست پایه‌ها
# =====================================================
def get_grades():
    try:
        school_id = get_current_school_id()
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT DISTINCT grade FROM students WHERE is_active = 1 AND school_id = ? ORDER BY grade",
            (school_id,)
        )
        rows = cursor.fetchall()
        conn.close()
        return jsonify({'success': True, 'grades': [row['grade'] for row in rows]})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


# =====================================================
# لیست کلاس‌ها
# =====================================================
def get_classes():
    try:
        school_id = get_current_school_id()
        grade = request.args.get('grade', '').strip()
        conn = get_db_connection()
        cursor = conn.cursor()
        if grade:
            cursor.execute(
                "SELECT DISTINCT class_name FROM students WHERE is_active = 1 AND school_id = ? AND grade = ? ORDER BY class_name",
                (school_id, grade)
            )
        else:
            cursor.execute(
                "SELECT DISTINCT class_name FROM students WHERE is_active = 1 AND school_id = ? ORDER BY class_name",
                (school_id,)
            )
        rows = cursor.fetchall()
        conn.close()
        return jsonify({'success': True, 'classes': [row['class_name'] for row in rows]})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


# =====================================================
# آپلود اکسل
# =====================================================
@admin_required
def upload_excel():
    try:
        from openpyxl import load_workbook
        import io

        school_id = get_current_school_id()

        if 'file' not in request.files:
            return jsonify({'success': False, 'error': 'فایلی انتخاب نشده'})

        file = request.files['file']
        if not file.filename:
            return jsonify({'success': False, 'error': 'فایلی انتخاب نشده'})

        if not file.filename.lower().endswith(('.xlsx', '.xlsm')):
            return jsonify({'success': False, 'error': 'فقط فایل xlsx پشتیبانی می‌شود'})

        file_bytes = file.read()
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active

        conn = get_db_connection()
        cursor = conn.cursor()

        success_count = 0
        error_count = 0
        errors = []

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[0]:
                continue

            try:
                national_code = str(row[0]).strip() if row[0] else ''
                first_name = str(row[1]).strip() if len(row) > 1 and row[1] else ''
                last_name = str(row[2]).strip() if len(row) > 2 and row[2] else ''
                father_name = str(row[3]).strip() if len(row) > 3 and row[3] else ''
                grade = str(row[4]).strip() if len(row) > 4 and row[4] else ''
                class_name = str(row[5]).strip() if len(row) > 5 and row[5] else ''
                field = str(row[6]).strip() if len(row) > 6 and row[6] else ''
                phone1 = str(row[7]).strip() if len(row) > 7 and row[7] else ''
                phone2 = str(row[8]).strip() if len(row) > 8 and row[8] else ''

                national_code = national_code.split('.')[0]
                phone1 = phone1.split('.')[0]
                phone2 = phone2.split('.')[0]

                if not national_code or not first_name or not last_name or not grade or not class_name:
                    errors.append(f"سطر {row_idx}: فیلدهای اجباری خالی")
                    error_count += 1
                    continue

                if len(national_code) != 10 or not national_code.isdigit():
                    errors.append(f"سطر {row_idx}: کد ملی نامعتبر")
                    error_count += 1
                    continue

                cursor.execute(
                    "SELECT id FROM students WHERE national_code = ? AND school_id = ?",
                    (national_code, school_id)
                )
                if cursor.fetchone():
                    errors.append(f"سطر {row_idx}: کد ملی {national_code} تکراری")
                    error_count += 1
                    continue

                cursor.execute('''
                    INSERT INTO students
                    (school_id, national_code, first_name, last_name, father_name, grade, class_name, field, phone1, phone2, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (school_id, national_code, first_name, last_name, father_name, grade, class_name, field, phone1, phone2, get_today_jalali()))

                success_count += 1
            except Exception as e:
                errors.append(f"سطر {row_idx}: {str(e)}")
                error_count += 1

        conn.commit()
        conn.close()

        return jsonify({
            'success': True,
            'message': f'{success_count} دانش‌آموز اضافه شد',
            'success_count': success_count,
            'error_count': error_count,
            'errors': errors[:20]
        })
    except Exception as e:
        logger.error(f"خطا در آپلود اکسل: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500
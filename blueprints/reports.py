from flask import render_template, request, jsonify, send_file, session
from core.database import get_db_connection
from core.logger import get_logger
from blueprints.decorators import admin_required
from datetime import datetime
import io

logger = get_logger(__name__)


def reports_page():
    return render_template('reports.html')


# =====================================================
# API: دریافت گزارش با فیلتر
# =====================================================
def get_report():
    try:
        date_from = request.args.get('date_from', '').strip()
        date_to = request.args.get('date_to', '').strip()
        grade = request.args.get('grade', '').strip()
        class_name = request.args.get('class_name', '').strip()
        status_filter = request.args.get('status', '').strip()

        query = '''
            SELECT 
                s.national_code, s.first_name, s.last_name,
                s.grade, s.class_name,
                a.attendance_date, a.status, a.note, a.recorded_by
            FROM attendance a
            JOIN students s ON a.student_id = s.id
            WHERE 1=1
        '''
        params = []

        if date_from:
            query += " AND a.attendance_date >= ?"
            params.append(date_from)
        if date_to:
            query += " AND a.attendance_date <= ?"
            params.append(date_to)
        if grade:
            query += " AND s.grade = ?"
            params.append(grade)
        if class_name:
            query += " AND s.class_name = ?"
            params.append(class_name)
        if status_filter:
            query += " AND a.status = ?"
            params.append(status_filter)

        query += " ORDER BY a.attendance_date DESC, s.grade, s.class_name, s.last_name"

        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(query, params)
        rows = cursor.fetchall()
        conn.close()

        # شمارش بر اساس وضعیت
        summary = {'present': 0, 'absent': 0, 'leave': 0, 'late': 0}
        records = []
        for r in rows:
            d = dict(r)
            records.append(d)
            if d['status'] in summary:
                summary[d['status']] += 1

        return jsonify({
            'success': True,
            'records': records,
            'count': len(records),
            'summary': summary
        })

    except Exception as e:
        logger.error(f"خطا در گزارش: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500


# =====================================================
# API: خروجی اکسل
# =====================================================
@admin_required
def export_excel():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill

        date_from = request.args.get('date_from', '').strip()
        date_to = request.args.get('date_to', '').strip()
        grade = request.args.get('grade', '').strip()
        class_name = request.args.get('class_name', '').strip()

        query = '''
            SELECT 
                s.national_code, s.first_name, s.last_name,
                s.grade, s.class_name,
                a.attendance_date, a.status, a.note, a.recorded_by
            FROM attendance a
            JOIN students s ON a.student_id = s.id
            WHERE 1=1
        '''
        params = []

        if date_from:
            query += " AND a.attendance_date >= ?"
            params.append(date_from)
        if date_to:
            query += " AND a.attendance_date <= ?"
            params.append(date_to)
        if grade:
            query += " AND s.grade = ?"
            params.append(grade)
        if class_name:
            query += " AND s.class_name = ?"
            params.append(class_name)

        query += " ORDER BY a.attendance_date DESC, s.grade, s.class_name, s.last_name"

        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(query, params)
        rows = cursor.fetchall()
        conn.close()

        # ساخت فایل اکسل
        wb = Workbook()
        ws = wb.active
        ws.title = "گزارش حضور و غیاب"
        ws.sheet_view.rightToLeft = True

        # هدر
        headers = ['کد ملی', 'نام', 'نام خانوادگی', 'پایه', 'کلاس', 'تاریخ', 'وضعیت', 'یادداشت', 'ثبت‌کننده']
        ws.append(headers)

        # استایل هدر
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="764ba2", end_color="764ba2", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # ترجمه وضعیت‌ها
        status_map = {
            'present': 'حاضر',
            'absent': 'غایب',
            'leave': 'مرخصی',
            'late': 'تاخیر'
        }

        # داده‌ها
        for row in rows:
            ws.append([
                row['national_code'],
                row['first_name'],
                row['last_name'],
                row['grade'],
                row['class_name'],
                row['attendance_date'],
                status_map.get(row['status'], row['status']),
                row['note'] or '',
                row['recorded_by'] or ''
            ])

        # تنظیم عرض ستون‌ها
        widths = [15, 15, 15, 10, 10, 15, 12, 20, 15]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w

        # ذخیره در حافظه
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"attendance_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )

    except Exception as e:
        logger.error(f"خطا در خروجی اکسل: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500